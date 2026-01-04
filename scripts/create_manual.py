from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import date
import os

ngc_path = "fraeser_schaerfen.ngc"
output_path = "Bedienungsanleitung_Grinder.docx"

# Read a short NGC excerpt
ngc_excerpt = []
try:
    with open(ngc_path, "r", encoding="utf-8") as f:
        for _ in range(30):
            line = f.readline()
            if not line:
                break
            ngc_excerpt.append(line.rstrip())
except Exception as e:
    ngc_excerpt = [f"(Fehler beim Lesen von {ngc_path}: {e})"]

# Document
doc = Document()

# Title
title = doc.add_paragraph()
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run = title.add_run("Bedienungsanleitung — Grinder")
run.bold = True
run.font.size = Pt(16)

doc.add_paragraph(f"Version: 1.0")
doc.add_paragraph(f"Datum: {date.today().isoformat()}")
doc.add_paragraph("")

# Inhaltsverzeichnis (Feld; in Word mit F9 aktualisieren)
doc.add_heading("Inhaltsverzeichnis", level=2)
# Ein echtes Inhaltsverzeichnis-Feld einfügen (Word aktualisiert es beim Öffnen/mit F9)
p = doc.add_paragraph()
run = p.add_run()
# TOC field: Word interpretiert dieses Feld und kann es beim Öffnen aktualisieren
fldBegin = OxmlElement('w:fldChar')
fldBegin.set(qn('w:fldCharType'), 'begin')
run._r.append(fldBegin)
instr = OxmlElement('w:instrText')
instr.set(qn('xml:space'), 'preserve')
instr.text = 'TOC \\o "1-3" \\h \\z \\u'
run._r.append(instr)
fldSeparate = OxmlElement('w:fldChar')
fldSeparate.set(qn('w:fldCharType'), 'separate')
run._r.append(fldSeparate)
fldEnd = OxmlElement('w:fldChar')
fldEnd.set(qn('w:fldCharType'), 'end')
run._r.append(fldEnd)

doc.add_paragraph("(Hinweis: Inhaltsverzeichnis in Word aktualisieren: markiere das Verzeichnis und drücke F9)")

doc.add_page_break()

# Sections
doc.add_heading("1. Einleitung", level=2)
doc.add_paragraph("Diese Bedienungsanleitung beschreibt die Verwendung des G-Code Generators 'Grinder.py' zur automatischen Erstellung von G-Code für das Fräser-Schärfen sowie Hinweise zur Nutzung der erzeugten Datei 'fraeser_schaerfen.ngc'. Zielgruppe sind Anwender mit Grundkenntnissen in CNC-Programmierung.")

doc.add_heading("2. Sicherheitshinweise", level=2)
doc.add_paragraph('- Vor dem Ausführen des G-Codes Maschinenbegrenzungen, Werkstückspannung und Werkzeuglänge prüfen.\n- Spindeldrehzahl und Vorschub prüfen bevor das Programm gestartet wird.\n- Erst einen "Trockenlauf" (kein Werkstück) mit reduzierter Geschwindigkeit durchführen.\n- Not-Aus-Funktion der Maschine kennen und testen.')

doc.add_heading("3. Voraussetzungen & Installation", level=2)
doc.add_paragraph("Voraussetzungen:")
doc.add_paragraph("- Python 3.8+\n- Die Datei 'Grinder.py' und die Konfigurationsdatei in 'ToolLib/template1.json' im selben Verzeichnis\n- Schreibrechte im Verzeichnis für die Ausgabedatei.")
doc.add_paragraph("Installation (Windows / macOS / Linux):")
doc.add_paragraph("1. Virtuelle Umgebung erstellen (optional): python -m venv .venv\n2. Abhängigkeiten installieren: pip install -r requirements.txt")

doc.add_heading('Linux Hinweise', level=3)
doc.add_paragraph('Für Debian/Ubuntu (Beispiel):')
doc.add_paragraph('- Systempakete: sudo apt update && sudo apt install python3-pip python3-tk build-essential libffi-dev libssl-dev python3-dev')
doc.add_paragraph("- Empfohlene Python-Pakete: pip install -r requirements.txt\n- Hinweis: Für Keyring mit GNOME empfiehlt sich 'secretstorage' (pip install secretstorage) und D-Bus auf dem System; unter KDE/Plasma kann KWallet verwendet werden.")
doc.add_paragraph('Hinweis zu LinuxCNC: Grinder kann als externes Werkzeug genutzt werden: Schreibe ein kleines Wrapper-Skript oder binde einen Button in der LinuxCNC‑GUI, der Grinder mit -c <config> aufruft. Für tiefere Integration ist ein genaueres Ziel-Setup nötig (LinuxCNC Version, GUI-Toolkit).')

doc.add_heading("4. Bedienung (Grinder.py)", level=2)
doc.add_paragraph("Der G-Code Generator erzeugt die Datei gemäß Konfiguration. Standard-Aufruf:")
cmd = doc.add_paragraph()
cmd.add_run("python Grinder.py -c ToolLib/template1.json").italic = True

doc.add_paragraph("Wichtige CLI-Optionen:")
doc.add_paragraph("- --edit : Öffnet nach Auswahl den JSON-Editor zum Bearbeiten/Speichern der Konfiguration.")
doc.add_paragraph("- --open <datei> : Öffnet eine beliebige JSON-Datei im Editor und beendet danach (praktisch zum schnellen Editieren).")
doc.add_paragraph("Editor-Funktionen:")
doc.add_paragraph("- Zeilennummern, Statusleiste mit JSON-Validität und Cursor-Position.")
doc.add_paragraph("- 'Öffnen...', 'Speichern', 'Speichern unter...' und 'Formatieren' (Prettify / JSON-Formatierung).")
doc.add_paragraph("- Tastaturkürzel: Ctrl+S = Speichern, Ctrl+F = Formatieren.")
doc.add_paragraph("Ablage des letzten Pfads:")
doc.add_paragraph("Die zuletzt verwendete Konfigurationsdatei wird jetzt im Projekt-Root unter '.grinder_last_config.json' gespeichert (Fallback auf Benutzer-Home).")

doc.add_paragraph("Erklärung wichtiger Konfigurationsparameter (ToolLib/template1.json):")
params = [
    ("maschine.safe_z", "Sicherheits-Z-Höhe (z.B. 20.0)"),
    ("maschine.start_x", "Start X-Position"),
    ("maschine.a_start/a_end/a_steps", "Start-/Endwinkel und Schritte für A-Achse"),
    ("schleifscheibe.anzahl_passe", "Anzahl der Schleifpässe"),
    ("schleifscheibe.zustellung_pro_pass", "Zustellung pro Pass (mm)"),
    ("werkzeug.schneidenanzahl", "Anzahl der Schneiden am Fräser"),
    ("werkzeug.schneidenlaenge", "Länge der zu schleifenden Schneide (mm)"),
    ("werkzeug.drall_grad_pro_mm", "Drall in Grad pro mm (oder drall_steigung = mm pro 360°)")
]
for k, v in params:
    p = doc.add_paragraph()
    p.add_run(k + ": ").bold = True
    p.add_run(v)

doc.add_heading("5. Bedienung (G-Code: fraeser_schaerfen.ngc)", level=2)
doc.add_paragraph("Die generierte G-Code Datei enthält Header, mehrere Schleifdurchgänge und ein Footer. Wichtige Blöcke:")
doc.add_paragraph("- Header: Maschinenmodus, Werkzeugwechsel, Spindel an\n- Schleifdurchgang: Z-Anfahrt, Schneidenweise Schleifen mit X/A Bewegungen\n- Footer: Z-Sicher, Spindel aus, Programm Ende (M30)")

doc.add_heading("Frontschärfen (FrontGrinder.py)", level=3)
doc.add_paragraph("Ein separates Hilfsprogramm `FrontGrinder.py` erzeugt G-Code speziell zum Schärfen der Stirnfläche des Fräsers. Konfiguriere den Abschnitt `front` in der JSON (siehe Template) und rufe z.B.: `python FrontGrinder.py -c ToolLib/template1.json --nogui` auf.")
doc.add_paragraph('Hinweis zum Ablaufschema: Bei jeder Zustellung bewegt sich die Z‑Achse von Z0 um den Tool‑Radius nach unten (Z = -ToolRadius + aktuelle Zustellung). Nach dem Vorschub in X‑Richtung zieht sich das Werkzeug in negative X‑Richtung um `front.retract_x` (Standard = maschine.safe_z) zurück, fährt auf Z0 und anschließend auf Z‑Sicherheitsabstand hoch. Die Start‑X‑Position kann pro Pass um `front.x_step` in Richtung `front.x_end` verschoben werden (optionale progressive Schritte).\n\nBeim Start des Programms wird optional ein Werkzeugwechsel (`Tn M6`) ausgeführt. Die **Spindel wird erst nach dem Anfahren der Sicherheitshöhe (`maschine.safe_z` bzw. `front.z_sicher`) gestartet** (`M3 S<drehzahl>`). Du kannst das Verhalten auf "Start nur beim ersten Pass" umschalten (z.B. `front.spindle_on_first_pass = true`). Zusätzlich kann die Dwell‑Zeit (Wartezeit nach Spindelstart) über `spindle_dwell_ms` konfiguriert werden.')

doc.add_paragraph("Beispiel (Auszug aus 'fraeser_schaerfen.ngc'):")
pre = doc.add_paragraph()
pre.style = 'Intense Quote'
for line in ngc_excerpt:
    pre.add_run(line + "\n")

# Beispiel: Dokumenterzeugung (Konfig-Beispiel & optionaler Screenshot)
doc.add_heading('Beispiel: Dokumenterzeugung', level=2)
# Versuche, Konfig-Werte aus ToolLib/template1.json zu lesen und in einer Tabelle darzustellen
template_path = os.path.join(os.path.dirname(__file__), '..', 'ToolLib', 'template1.json')
try:
    import json
    with open(template_path, 'r', encoding='utf-8') as f:
        cfg = json.load(f)
        m = cfg.get('maschine', {})
        s = cfg.get('schleifscheibe', {})
        w = cfg.get('werkzeug', {})
        table = doc.add_table(rows=1, cols=2)
        hdr = table.rows[0].cells
        hdr[0].text = 'Parameter'
        hdr[1].text = 'Wert'
        # Füge einige typische Parameter hinzu
        for k in ('safe_z', 'start_x'):
            row = table.add_row().cells
            row[0].text = k
            row[1].text = str(m.get(k, ''))
        row = table.add_row().cells
        row[0].text = 'anzahl_passe'
        row[1].text = str(s.get('anzahl_passe', ''))
        row = table.add_row().cells
        row[0].text = 'schneidenanzahl'
        row[1].text = str(w.get('schneidenanzahl', ''))
except Exception:
    doc.add_paragraph('(Template nicht gefunden für Beispiel-Tabelle)')

# Optional: Screenshot einfügen, wenn scripts/screenshot.png existiert
img_path = os.path.join(os.path.dirname(__file__), 'screenshot.png')
doc.add_heading('Screenshot (optional)', level=3)
if os.path.exists(img_path):
    try:
        doc.add_picture(img_path, width=Inches(4))
    except Exception as e:
        doc.add_paragraph(f'(Bild konnte nicht eingefügt werden: {e})')
else:
    doc.add_paragraph('(Kein Bild gefunden. Legen Sie optional scripts/screenshot.png ab, um ein Bild einzufügen.)')

doc.add_heading("6. Wartung & Pflege", level=2)
doc.add_paragraph("- Regelmäßig die Schleifscheibe prüfen und bei Verschleiß ersetzen.\n- Achsenführungen schmieren und Maschinenparameter kontrollieren.\n- Backup der Konfigurationsdatei 'ToolLib/template1.json' aufbewahren.")

doc.add_heading("7. Fehlerbehebung", level=2)
doc.add_paragraph("Häufige Probleme und Lösungen:")
doc.add_paragraph("- Keine Ausgabedatei erzeugt: Prüfen ob 'ausgabe.datei' in der Konfig gesetzt ist und Schreibrechte bestehen.\n- Unerwartete Achsenbewegungen: Konfig Werte (start_x, safe_z, a_start) überprüfen und zuerst mit Trockenlauf testen.\n- Spindel dreht mit falscher Drehzahl: 'schleifscheibe.drehzahl' prüfen.")


doc.add_heading("8. Anhang", level=2)
doc.add_paragraph("Dateien im Projekt:\n- Grinder.py (G-Code Generator)\n- ToolLib/template1.json (Konfiguration)\n- fraeser_schaerfen.ngc (Beispielausgabe)")

doc.add_heading("9. Release bauen (optional)", level=2)
doc.add_paragraph("Für Windows-Exe: Das Projekt enthält ein Hilfsskript 'scripts/build_release.py', das PyInstaller aufruft und eine ZIP-Datei mit der EXE und einer README erzeugt. Beispiel: \npython scripts/build_release.py --version 1.0.0\nOptional: --icon path/to/icon.ico\nHinweis: PyInstaller muss installiert sein (pip install pyinstaller).")

doc.add_heading("10. Upload-Funktion (FTP / SFTP)", level=2)
doc.add_paragraph("Das Editor-Fenster bietet jetzt: 'Server Verwaltung' (Hinzufügen / Bearbeiten / Löschen), Testverbindung, und beim Erzeugen des G-Codes wird angeboten, die Datei auf einen konfigurierten Server hochzuladen. Unterstützt werden FTP (ftplib) und SFTP (paramiko).")
doc.add_paragraph('Server-Passwörter können optional sicher im System-Keyring gespeichert werden (Windows Credential Manager, macOS Keychain, etc.). Beim Hinzufügen/Bearbeiten gibt es eine Option "Password im System-Keyring speichern".')
doc.add_paragraph("Server werden in '.grinder_servers.json' im Projekt-Root gespeichert. Eintrag: {name, type(ftp|sftp), host, port, username, password, remote_path, key_path}")


doc.add_paragraph("")

doc.save(output_path)
print(f"Bedienungsanleitung erstellt: {output_path}")
