"""
Kurzes Beispiel: zeigt die wichtigsten python-docx Operationen
- Überschrift, formatierter Text, Liste, Tabelle
- Kopfzeile, Fußzeile
- Einfügen eines NGC-Auszugs (fraeser_schaerfen.ngc)
- Optional: Einfügen eines Bildes (wenn vorhanden)

Ausführen: python scripts/example_docx.py
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from datetime import date
import os
import json

HERE = os.path.dirname(__file__)
NGC = os.path.join(HERE, '..', 'fraeser_schaerfen.ngc')
TEMPLATE = os.path.join(HERE, '..', 'ToolLib', 'template1.json')
OUT = os.path.join(HERE, 'example_docx_demo.docx')
IMG = os.path.join(HERE, 'screenshot.png')  # optional

# Lese ein paar Zeilen aus der NGC-Datei
ngc_excerpt = []
try:
    with open(NGC, 'r', encoding='utf-8') as f:
        for _ in range(20):
            line = f.readline()
            if not line:
                break
            ngc_excerpt.append(line.rstrip())
except Exception:
    ngc_excerpt = ['(NGC-Datei nicht gefunden oder konnte nicht gelesen werden)']

# Versuch: einige Konfigwerte aus template1.json lesen (falls vorhanden)
config_vals = {}
try:
    with open(TEMPLATE, 'r', encoding='utf-8') as f:
        cfg = json.load(f)
        m = cfg.get('maschine', {})
        s = cfg.get('schleifscheibe', {})
        w = cfg.get('werkzeug', {})
        config_vals = {
            'safe_z': m.get('safe_z'),
            'start_x': m.get('start_x'),
            'anzahl_passe': s.get('anzahl_passe'),
            'schneidenanzahl': w.get('schneidenanzahl')
        }
except Exception:
    config_vals = {}

# Document anlegen
doc = Document()

# Kopf: Titel
h = doc.add_paragraph()
h.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
r = h.add_run('Beispiel: python-docx Demonstration')
r.bold = True
r.font.size = Pt(14)

doc.add_paragraph(f"Datum: {date.today().isoformat()}")
doc.add_paragraph('')

# Formatierter Absatz
p = doc.add_paragraph()
p.add_run('Dieses Dokument zeigt: ').bold = True
p.add_run('Überschriften, fetten Text, Listen, Tabellen, Einfügung eines NGC-Auszugs und optionales Bild.').italic = True

# Liste
doc.add_paragraph('Schritte:', style='List Bullet')
doc.add_paragraph('1) Generiere G-Code mit Grinder.py', style='List Number')
doc.add_paragraph('2) Führe Trockenlauf aus (kein Werkstück)', style='List Number')

# Tabelle mit Konfig-Werten (falls vorhanden)
doc.add_heading('Konfigurationswerte (aus ToolLib/template1.json)', level=2)
if config_vals:
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = 'Parameter'
    hdr[1].text = 'Wert'
    for k, v in config_vals.items():
        row = table.add_row().cells
        row[0].text = str(k)
        row[1].text = str(v)
else:
    doc.add_paragraph('(Konfigurationsdatei nicht gefunden oder leer)')

# NGC-Auszug
doc.add_heading('Auszug: fraeser_schaerfen.ngc', level=2)
ngc_para = doc.add_paragraph()
ngc_para.style = 'Intense Quote'
for line in ngc_excerpt:
    ngc_para.add_run(line + '\n')

# Optional: Bild einfügen (wenn vorhanden)
doc.add_heading('Screenshot (optional)', level=2)
if os.path.exists(IMG):
    try:
        doc.add_picture(IMG, width=Inches(4))
    except Exception as e:
        doc.add_paragraph(f'(Bild konnte nicht eingefügt werden: {e})')
else:
    doc.add_paragraph('(Kein Bild gefunden. Legen Sie optional scripts/screenshot.png ab um es einzufügen.)')

# Kopfzeile und Fußzeile
section = doc.sections[0]
header = section.header
header_para = header.paragraphs[0]
header_para.text = 'Grinder — Beispiel-Dokument'
footer = section.footer
footer_para = footer.paragraphs[0]
footer_para.text = 'Erzeugt mit python-docx'

# Speichern
doc.save(OUT)
print(f'Beispieldokument erstellt: {OUT}')
